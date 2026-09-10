import openpyxl

# Column headers differ slightly between the 1,000-row and 10,000-row databases.
# Map each known header to a common internal field name.
_HEADER_MAP = {
    "video id": "id",
    "id": "id",
    "content pillar / category": "pillar",
    "pillar category": "pillar",
    "hook (0-3s)": "hook",
    "script body (voiceover)": "script_body",
    "script body (narrative voiceover)": "script_body",
    "call to action (cta)": "cta",
    "full script text": "full_text",
    "featured character(s)": "character",
    "primary character": "character",
    "character visual description": "character_visual",
    "character visual details": "character_visual",
    "visual sequence & animation cue": "animation_cue",
    "visual animation & camera cue": "animation_cue",
    "audio / music style": "audio_style",
    "audio & sfx direction": "audio_style",
}


def _find_script_sheet(workbook):
    for sheet in workbook.worksheets:
        header = [c.value for c in sheet[1]]
        if any(isinstance(h, str) and h.strip().lower() in ("video id", "id") for h in header):
            return sheet
    raise ValueError("No script sheet found (expected a header row with 'Video ID' or 'ID').")


def load_row(xlsx_path: str, video_id: str) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    sheet = _find_script_sheet(wb)

    header = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    fields = [_HEADER_MAP.get(h.strip().lower(), h) if isinstance(h, str) else h for h in header]

    for row in sheet.iter_rows(min_row=2):
        values = [c.value for c in row]
        record = dict(zip(fields, values))
        if str(record.get("id", "")).strip() == video_id.strip():
            return _normalize(record)

    raise KeyError(f"Video ID '{video_id}' not found in {xlsx_path}")


def iter_rows(xlsx_path: str, start_id: str = None, count: int = None):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    sheet = _find_script_sheet(wb)

    header = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    fields = [_HEADER_MAP.get(h.strip().lower(), h) if isinstance(h, str) else h for h in header]

    started = start_id is None
    yielded = 0
    for row in sheet.iter_rows(min_row=2):
        if count is not None and yielded >= count:
            return
        values = [c.value for c in row]
        record = dict(zip(fields, values))
        row_id = str(record.get("id", "")).strip()
        if not row_id:
            continue
        if not started:
            if row_id == start_id.strip():
                started = True
            else:
                continue
        yield _normalize(record)
        yielded += 1


def _normalize(record: dict) -> dict:
    full_text = record.get("full_text")
    if not full_text:
        parts = [record.get("hook"), record.get("script_body"), record.get("cta")]
        full_text = " ".join(p for p in parts if p)
    return {
        "id": str(record.get("id", "")).strip(),
        "pillar": record.get("pillar") or "",
        "hook": record.get("hook") or "",
        "script_body": record.get("script_body") or "",
        "cta": record.get("cta") or "",
        "full_text": full_text or "",
        "character": record.get("character") or "",
        "character_visual": record.get("character_visual") or "",
        "animation_cue": record.get("animation_cue") or "",
        "audio_style": record.get("audio_style") or "",
    }
